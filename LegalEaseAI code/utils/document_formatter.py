from io import BytesIO
from pathlib import Path
from typing import Optional
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from fpdf import FPDF


def sanitize_text(text: str) -> str:

    replacements = {

        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u00a0": " "
    }


    for old, new in replacements.items():

        text = text.replace(old, new)


    return text.strip()


def get_lines(text):

    text = sanitize_text(text)

    return [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]


# --------------------------------------------------
# DOCX
# --------------------------------------------------

def format_docx(
    text: str,
    doc_type: str,
    logo_bytes: Optional[bytes] = None
):

    document = Document()


    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)


    # Normal font

    normal_style = document.styles["Normal"]

    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11)


    # Logo

    if logo_bytes:

        paragraph = document.add_paragraph()

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run()

        run.add_picture(
            BytesIO(logo_bytes),
            width=Inches(1.25)
        )


    # Title

    title = document.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title.add_run(
        doc_type.upper()
    )

    title_run.bold = True

    title_run.font.name = "Times New Roman"

    title_run.font.size = Pt(16)


    # Document content

    for line in get_lines(text):

        # Numbered heading

        if re.match(
            r"^(\d+[\.\)]|SECTION\s+\d+)",
            line,
            re.IGNORECASE
        ):

            paragraph = document.add_paragraph()

            run = paragraph.add_run(line)

            run.bold = True


        # Bullet

        elif line.startswith(
            ("-", "•", "*")
        ):

            document.add_paragraph(
                line.lstrip("-•* ").strip(),
                style="List Bullet"
            )


        # Normal paragraph

        else:

            document.add_paragraph(line)


    # Footer

    footer = section.footer.paragraphs[0]

    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer_run = footer.add_run(
        "LegalEase - AI-assisted draft | "
        "Review with a qualified legal professional"
    )

    footer_run.italic = True


    output = BytesIO()

    document.save(output)

    return output.getvalue()


# --------------------------------------------------
# PDF
# --------------------------------------------------

class LegalPDF(FPDF):

    def __init__(self, logo_path=None):

        super().__init__()

        self.logo_path = logo_path

        self.set_auto_page_break(
            auto=True,
            margin=18
        )


    def header(self):

        if (
            self.logo_path
            and Path(self.logo_path).exists()
        ):

            try:

                self.image(
                    self.logo_path,
                    x=90,
                    y=8,
                    w=30
                )

                self.ln(18)

            except Exception:

                self.ln(4)

        else:

            self.ln(4)


    def footer(self):

        self.set_y(-14)

        self.set_font(
            "Helvetica",
            "I",
            8
        )

        self.cell(
            0,
            8,
            "LegalEase - AI-assisted draft | "
            "Review with a qualified legal professional",
            align="C"
        )


def format_pdf(
    text: str,
    doc_type: str,
    logo_bytes: Optional[bytes] = None
):

    temporary_logo = None


    if logo_bytes:

        import tempfile

        temp = tempfile.NamedTemporaryFile(
            delete=False,
            suffix=".png"
        )

        temp.write(logo_bytes)

        temp.close()

        temporary_logo = temp.name


    pdf = LegalPDF(
        temporary_logo
    )


    pdf.add_page()


    pdf.set_margins(
        18,
        18,
        18
    )


    # Document title

    pdf.set_font(
        "Helvetica",
        "B",
        15
    )

    pdf.cell(
        0,
        10,
        sanitize_text(
            doc_type
        ).upper(),
        new_x="LMARGIN",
        new_y="NEXT",
        align="C"
    )


    pdf.ln(4)


    # Content

    for line in get_lines(text):

        if re.match(
            r"^(\d+[\.\)]|SECTION\s+\d+)",
            line,
            re.IGNORECASE
        ):

            pdf.set_font(
                "Helvetica",
                "B",
                11
            )

            pdf.multi_cell(
                0,
                6,
                line
            )


        elif line.startswith(
            ("-", "•", "*")
        ):

            pdf.set_font(
                "Helvetica",
                "",
                10.5
            )

            pdf.multi_cell(
                0,
                6,
                "- "
                + line.lstrip(
                    "-•* "
                ).strip()
            )


        else:

            pdf.set_font(
                "Helvetica",
                "",
                10.5
            )

            pdf.multi_cell(
                0,
                6,
                line
            )


        pdf.ln(1)


    pdf_data = bytes(
        pdf.output(dest="S")
    )


    # Delete temporary logo

    if temporary_logo:

        try:

            Path(
                temporary_logo
            ).unlink()

        except OSError:

            pass


    return pdf_data